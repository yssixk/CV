"""Security layer tests — the acceptance criteria from the build brief live here."""
from __future__ import annotations

import io
import time
import zipfile

import pytest

from cv_analyzer.security.rate_limit import RateLimitExceeded, RateLimiter, enforce
from cv_analyzer.security.redact import redact_text
from cv_analyzer.security.validate_upload import validate_upload


# ---------------------------------------------------------------------------
# upload validation
# ---------------------------------------------------------------------------

def test_oversized_file_rejected_before_parsing():
    data = b"x" * (5 * 1024 * 1024 + 10)
    report = validate_upload(data, "cv.txt", max_mb=5.0)
    assert not report.ok
    assert "size cap" in report.error_message


def test_txt_renamed_to_pdf_is_caught():
    """Required acceptance test: mismatched extension is caught by magic bytes."""
    data = b"hello world, this is definitely not a pdf document"
    report = validate_upload(data, "cv.pdf", max_mb=5.0)
    assert not report.ok
    assert "renamed" in report.error_message or "signature" in report.error_message


def test_exe_bytes_with_pdf_extension_rejected():
    data = b"MZ\x90\x00" + b"\x00" * 64      # PE executable magic
    report = validate_upload(data, "cv.pdf", max_mb=5.0)
    assert not report.ok


def test_disallowed_extension_rejected():
    report = validate_upload(b"hello", "cv.exe", max_mb=5.0)
    assert not report.ok
    assert "not allowed" in report.error_message


def test_empty_file_rejected():
    assert not validate_upload(b"", "cv.txt", max_mb=5.0).ok


def test_valid_text_file_passes():
    report = validate_upload("BUDI SANTOSO\nSUMMARY\nhard worker".encode(), "cv.txt", max_mb=5.0)
    assert report.ok and report.kind == "text"


def test_valid_pdf_signature_passes_validation():
    data = b"%PDF-1.4\n%fake but signed\n..."
    report = validate_upload(data, "cv.pdf", max_mb=5.0)
    assert report.ok and report.kind == "pdf"


def test_text_content_with_pdf_extension_rejected():
    report = validate_upload(b"just some text", "cv.pdf", max_mb=5.0)
    assert not report.ok


def test_docx_zip_bomb_guard():
    """A docx whose declared uncompressed size exceeds the cap is rejected."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("word/document.xml", b"0" * (10 * 1024 * 1024))   # 10 MB uncompressed
    report = validate_upload(buf.getvalue(), "cv.docx", max_mb=1.0)
    assert not report.ok
    assert "exceeds cap" in report.error_message or "zip bomb" in report.error_message


def test_valid_docx_passes_validation():
    import docx as docx_module

    buf = io.BytesIO()
    document = docx_module.Document()
    document.add_paragraph("BUDI SANTOSO")
    document.add_paragraph("SUMMARY")
    document.add_paragraph("Hard worker with Python skills.")
    document.save(buf)
    report = validate_upload(buf.getvalue(), "cv.docx", max_mb=5.0)
    assert report.ok and report.kind == "docx"


# ---------------------------------------------------------------------------
# redaction
# ---------------------------------------------------------------------------

def test_redacts_email():
    out = redact_text("Contact me at budi.santoso@example.com anytime.").redacted_text
    assert "budi.santoso@example.com" not in out
    assert "[EMAIL]" in out


def test_redacts_indonesian_phone():
    out = redact_text("Phone: +62 812-3456-7890").redacted_text
    assert "812-3456-7890" not in out
    assert "[PHONE]" in out


def test_redacts_url():
    out = redact_text("Portfolio at https://budisantoso.dev and linkedin.com/in/budi").redacted_text
    assert "budisantoso.dev" not in out


def test_redacts_name_line():
    out = redact_text("Budi Santoso\nJakarta, Indonesia").redacted_text
    assert out.splitlines()[0] == "[NAME]"


def test_redaction_preserves_non_pii_text():
    original = "Reduced report generation time by 40% with Python."
    assert redact_text(original).redacted_text == original


# ---------------------------------------------------------------------------
# rate limiting
# ---------------------------------------------------------------------------

def test_rate_limiter_blocks_after_cap():
    limiter = RateLimiter(max_calls=2, per_seconds=60.0)
    assert limiter.check("s1")[0] is True
    assert limiter.check("s1")[0] is True
    allowed, retry_after = limiter.check("s1")
    assert not allowed and retry_after > 0


def test_rate_limiter_window_slides():
    limiter = RateLimiter(max_calls=1, per_seconds=0.05)
    assert limiter.check("s2")[0] is True
    assert limiter.check("s2")[0] is False
    time.sleep(0.06)
    assert limiter.check("s2")[0] is True


def test_rate_limiter_is_per_session():
    limiter = RateLimiter(max_calls=1, per_seconds=60.0)
    assert limiter.check("visitor-a")[0] is True
    assert limiter.check("visitor-b")[0] is True   # isolated sessions


def test_enforce_raises():
    limiter = RateLimiter(max_calls=1, per_seconds=60.0)
    enforce(limiter, "s9")
    with pytest.raises(RateLimitExceeded):
        enforce(limiter, "s9")
